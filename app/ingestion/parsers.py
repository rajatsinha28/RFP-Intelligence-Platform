from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from PIL import Image
from pypdf import PdfReader
from readability import Document as ReadabilityDocument

from app.core.exceptions import DocumentParseError, UnsupportedDocumentError
from app.ingestion.types import ParsedDocument, ParsedElement
from app.observability.logging import get_logger
from app.security.prompts import sanitize_document_text

logger = get_logger()
PDF_PAGE_LABEL = re.compile(r"^\s*(\d{1,5})\s*$")
LINE_BREAK_HYPHEN = re.compile(r"(?<=[a-z])-[ \t]*\n[ \t]*(?=[a-z])")
SINGLE_LINE_BREAK = re.compile(r"(?<=\S)[ \t]*\n[ \t]*(?=\S)")


def extract_pdf_page_label(text: str, physical_page_number: int) -> tuple[int, str]:
    """Return a visible numeric page label when the PDF text exposes one.

    Many books include covers or front matter, so the reader's printed page label
    differs from the physical PDF page index. Only an isolated first or last line
    reasonably close to the physical index is treated as a label. The physical
    index is retained separately in element metadata.
    """

    lines = text.splitlines()
    nonempty = [index for index, line in enumerate(lines) if line.strip()]
    if not nonempty:
        return physical_page_number, text
    for line_index in (nonempty[0], nonempty[-1]):
        match = PDF_PAGE_LABEL.fullmatch(lines[line_index])
        if not match:
            continue
        label = int(match.group(1))
        if label > 0 and abs(label - physical_page_number) <= 100:
            remaining = lines[:line_index] + lines[line_index + 1 :]
            return label, "\n".join(remaining).strip()
    return physical_page_number, text


def clean_extracted_text(text: str) -> str:
    """Repair extraction artifacts without changing intentional compound words."""

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = LINE_BREAK_HYPHEN.sub("", text)
    paragraphs = [
        SINGLE_LINE_BREAK.sub(" ", paragraph)
        for paragraph in re.split(r"\n\s*\n", text)
        if paragraph.strip()
    ]
    return sanitize_document_text("\n\n".join(paragraphs))


def remove_repeated_page_edges(page_texts: list[str]) -> list[str]:
    """Remove headers/footers repeated on most pages while retaining page bodies."""

    if len(page_texts) < 3:
        return page_texts

    def lines(value: str) -> list[str]:
        return [line.strip() for line in value.splitlines() if line.strip()]

    def signature(value: str) -> str:
        return re.sub(r"\d+", "#", " ".join(value.lower().split()))

    edges: list[tuple[str | None, str | None]] = []
    counts: dict[str, int] = {}
    for text in page_texts:
        page_lines = lines(text)
        first = page_lines[0] if page_lines else None
        last = page_lines[-1] if page_lines else None
        edges.append((first, last))
        for value in {first, last} - {None}:
            key = signature(str(value))
            if len(key) >= 4:
                counts[key] = counts.get(key, 0) + 1
    threshold = max(2, int(len(page_texts) * 0.6 + 0.999))
    repeated = {key for key, count in counts.items() if count >= threshold}
    cleaned: list[str] = []
    for text, (first, last) in zip(page_texts, edges, strict=True):
        page_lines = text.splitlines()
        if first and signature(first) in repeated:
            first_index = next((i for i, line in enumerate(page_lines) if line.strip()), None)
            if first_index is not None:
                page_lines.pop(first_index)
        if last and signature(last) in repeated:
            last_index = next(
                (i for i in range(len(page_lines) - 1, -1, -1) if page_lines[i].strip()), None
            )
            if last_index is not None:
                page_lines.pop(last_index)
        cleaned.append("\n".join(page_lines).strip())
    return cleaned


def table_to_markdown(headers: list[str], rows: list[list[str]]) -> str:
    width = max([len(headers), *(len(row) for row in rows)], default=0)
    if width == 0:
        return ""
    normalized_headers = (headers + [f"Column {i + 1}" for i in range(len(headers), width)])[:width]
    normalized_rows = [(row + [""] * width)[:width] for row in rows]

    def clean(value: str) -> str:
        return " ".join(str(value).replace("|", "\\|").split())

    lines = [
        "| " + " | ".join(clean(value) for value in normalized_headers) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    lines.extend("| " + " | ".join(clean(value) for value in row) + " |" for row in normalized_rows)
    return "\n".join(lines)


class DocumentParser(ABC):
    extensions: tuple[str, ...]

    @abstractmethod
    def parse(self, path: Path, mime_type: str, source_url: str | None = None) -> ParsedDocument:
        raise NotImplementedError


class TextParser(DocumentParser):
    extensions = (".txt", ".md")

    def parse(self, path: Path, mime_type: str, source_url: str | None = None) -> ParsedDocument:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        elements: list[ParsedElement] = []
        current_heading: str | None = None
        buffer: list[str] = []

        def flush() -> None:
            if buffer:
                content = sanitize_document_text("\n".join(buffer))
                if content:
                    elements.append(ParsedElement(text=content, section=current_heading))
                buffer.clear()

        for line in text.splitlines():
            if re.match(r"^#{1,6}\s+", line):
                flush()
                current_heading = re.sub(r"^#{1,6}\s+", "", line).strip()
                elements.append(
                    ParsedElement(
                        text=current_heading, content_type="heading", heading=current_heading
                    )
                )
            elif line.strip():
                buffer.append(line)
            else:
                flush()
        flush()
        if not elements and text.strip():
            elements.append(ParsedElement(text=sanitize_document_text(text)))
        return ParsedDocument(path.stem, path.name, "text", mime_type, elements)


class JSONParser(DocumentParser):
    extensions = (".json",)

    def parse(self, path: Path, mime_type: str, source_url: str | None = None) -> ParsedDocument:
        try:
            data = json.loads(path.read_text(encoding="utf-8-sig"))
        except (json.JSONDecodeError, UnicodeDecodeError) as exc:
            raise DocumentParseError(f"Invalid JSON document: {exc}") from exc
        elements: list[ParsedElement] = []

        def visit(value: Any, field_path: str) -> None:
            if isinstance(value, dict):
                for key, child in value.items():
                    visit(child, f"{field_path}.{key}" if field_path else str(key))
            elif isinstance(value, list):
                if value and all(isinstance(item, dict) for item in value):
                    headers = sorted({key for item in value for key in item})
                    rows = [[str(item.get(key, "")) for key in headers] for item in value]
                    elements.append(
                        ParsedElement(
                            text=table_to_markdown(headers, rows),
                            content_type="table",
                            section=field_path,
                            table_id=f"json-{len(elements)}",
                            metadata={"json_path": field_path, "headers": headers},
                        )
                    )
                else:
                    for index, child in enumerate(value):
                        visit(child, f"{field_path}[{index}]")
            else:
                elements.append(
                    ParsedElement(
                        text=f"{field_path}: {value}",
                        section=field_path.rsplit(".", 1)[0] if "." in field_path else None,
                        metadata={"json_path": field_path},
                    )
                )

        visit(data, "")
        return ParsedDocument(path.stem, path.name, "json", mime_type, elements)


class CSVParser(DocumentParser):
    extensions = (".csv",)

    def parse(self, path: Path, mime_type: str, source_url: str | None = None) -> ParsedDocument:
        raw = path.read_text(encoding="utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(raw[:8192])
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(raw), dialect)
        all_rows = list(reader)
        if not all_rows:
            return ParsedDocument(path.stem, path.name, "csv", mime_type, [])
        headers = all_rows[0]
        rows = all_rows[1:]
        elements: list[ParsedElement] = []
        for offset in range(0, len(rows), 50):
            batch = rows[offset : offset + 50]
            elements.append(
                ParsedElement(
                    text=table_to_markdown(headers, batch),
                    content_type="table",
                    section=f"Rows {offset + 2}-{offset + len(batch) + 1}",
                    table_id=f"csv-{offset // 50}",
                    metadata={"headers": headers, "row_start": offset + 2},
                )
            )
        return ParsedDocument(
            path.stem,
            path.name,
            "csv",
            mime_type,
            elements,
            {"columns": headers, "row_count": len(rows)},
        )


class HTMLParser(DocumentParser):
    extensions = (".html", ".htm")

    def parse(self, path: Path, mime_type: str, source_url: str | None = None) -> ParsedDocument:
        raw = path.read_text(encoding="utf-8", errors="replace")
        readable = ReadabilityDocument(raw)
        soup = BeautifulSoup(readable.summary(), "lxml")
        original = BeautifulSoup(raw, "lxml")
        title = readable.short_title() or (original.title.string if original.title else path.stem)
        content_root = original.find("article") or original.find("main") or soup
        elements: list[ParsedElement] = []
        section: str | None = None
        table_index = 0
        for node in content_root.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "table"]):
            if node.name in {"h1", "h2", "h3", "h4"}:
                section = " ".join(node.get_text(" ", strip=True).split())
                if section:
                    elements.append(
                        ParsedElement(
                            text=section, content_type="heading", heading=section, section=section
                        )
                    )
            elif node.name == "table":
                rows = [
                    [
                        " ".join(cell.get_text(" ", strip=True).split())
                        for cell in row.find_all(["th", "td"])
                    ]
                    for row in node.find_all("tr")
                ]
                rows = [row for row in rows if row]
                if rows:
                    headers = rows[0]
                    elements.append(
                        ParsedElement(
                            text=table_to_markdown(headers, rows[1:]),
                            content_type="table",
                            section=section,
                            table_id=f"html-{table_index}",
                            metadata={"headers": headers},
                        )
                    )
                    table_index += 1
            else:
                text = " ".join(node.get_text(" ", strip=True).split())
                if text:
                    elements.append(
                        ParsedElement(text=sanitize_document_text(text), section=section)
                    )
        return ParsedDocument(
            str(title).strip(),
            path.name,
            "web" if source_url else "html",
            mime_type,
            elements,
            {"source_url": source_url},
        )


class ImageParser(DocumentParser):
    extensions = (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp")

    def __init__(self, languages: str = "eng") -> None:
        self.languages = languages

    def parse(self, path: Path, mime_type: str, source_url: str | None = None) -> ParsedDocument:
        try:
            import pytesseract

            image = Image.open(path)
            text = pytesseract.image_to_string(image, lang=self.languages)
        except Exception as exc:
            raise DocumentParseError(f"OCR failed for {path.name}: {exc}") from exc
        elements = (
            [
                ParsedElement(
                    text=sanitize_document_text(text),
                    content_type="image_ocr",
                    image_id="image-0",
                    metadata={"width": image.width, "height": image.height},
                )
            ]
            if text.strip()
            else []
        )
        return ParsedDocument(path.stem, path.name, "image", mime_type, elements)


class DOCXParser(DocumentParser):
    extensions = (".docx",)

    def __init__(self, languages: str = "eng") -> None:
        self.languages = languages

    def parse(self, path: Path, mime_type: str, source_url: str | None = None) -> ParsedDocument:
        try:
            document = DocxDocument(str(path))
        except Exception as exc:
            raise DocumentParseError(f"Invalid DOCX document: {exc}") from exc
        title = path.stem
        section: str | None = None
        elements: list[ParsedElement] = []
        for paragraph in document.paragraphs:
            text = " ".join(paragraph.text.split())
            if not text:
                continue
            style = paragraph.style.name if paragraph.style else ""
            if style.lower().startswith("title"):
                title = text
                elements.append(ParsedElement(text=text, content_type="heading", heading=text))
            elif style.lower().startswith("heading"):
                section = text
                elements.append(
                    ParsedElement(text=text, content_type="heading", heading=text, section=text)
                )
            else:
                elements.append(ParsedElement(text=sanitize_document_text(text), section=section))
        for index, table in enumerate(document.tables):
            rows = [[" ".join(cell.text.split()) for cell in row.cells] for row in table.rows]
            if rows:
                elements.append(
                    ParsedElement(
                        text=table_to_markdown(rows[0], rows[1:]),
                        content_type="table",
                        section=section,
                        table_id=f"docx-{index}",
                        metadata={"headers": rows[0]},
                    )
                )

        # Extract and OCR embedded images without mutating the source archive.
        try:
            import pytesseract

            with zipfile.ZipFile(path) as archive:
                media = [name for name in archive.namelist() if name.startswith("word/media/")]
                for index, name in enumerate(media):
                    with Image.open(io.BytesIO(archive.read(name))) as image:
                        text = pytesseract.image_to_string(image, lang=self.languages).strip()
                        if text:
                            elements.append(
                                ParsedElement(
                                    text=sanitize_document_text(text),
                                    content_type="image_ocr",
                                    section=section,
                                    image_id=f"docx-image-{index}",
                                )
                            )
        except (zipfile.BadZipFile, OSError):
            pass
        return ParsedDocument(title, path.name, "docx", mime_type, elements)


class PDFParser(DocumentParser):
    extensions = (".pdf",)

    def __init__(self, languages: str = "eng") -> None:
        self.languages = languages

    def parse(self, path: Path, mime_type: str, source_url: str | None = None) -> ParsedDocument:
        try:
            reader = PdfReader(path)
        except Exception as exc:
            raise DocumentParseError(f"Invalid PDF document: {exc}") from exc
        title = (
            str(reader.metadata.title) if reader.metadata and reader.metadata.title else path.stem
        )
        elements: list[ParsedElement] = []
        sparse_pages: list[int] = []
        page_labels: dict[int, int] = {}
        extracted_pages: list[tuple[int, str]] = []
        for index, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            text = sanitize_document_text(text)
            page_label, text = extract_pdf_page_label(text, index)
            page_labels[index] = page_label
            extracted_pages.append((page_label, text))
        cleaned_pages = remove_repeated_page_edges([text for _, text in extracted_pages])
        for index, ((page_label, _), text) in enumerate(
            zip(extracted_pages, cleaned_pages, strict=True), start=1
        ):
            text = clean_extracted_text(text)
            if len(text.strip()) < 40:
                sparse_pages.append(index)
            elif text.strip():
                elements.append(
                    ParsedElement(
                        text=text,
                        page_number=page_label,
                        section=f"Page {page_label}",
                        metadata={"pdf_page_number": index},
                    )
                )

        try:
            import pdfplumber

            with pdfplumber.open(path) as pdf:
                for page_number, pdf_page in enumerate(pdf.pages, start=1):
                    page_label = page_labels.get(page_number, page_number)
                    for table_index, table in enumerate(pdf_page.extract_tables() or []):
                        cleaned = [
                            [str(cell or "").strip() for cell in row] for row in table if row
                        ]
                        if cleaned:
                            elements.append(
                                ParsedElement(
                                    text=table_to_markdown(cleaned[0], cleaned[1:]),
                                    content_type="table",
                                    page_number=page_label,
                                    section=f"Page {page_label}",
                                    table_id=f"pdf-{page_number}-{table_index}",
                                    metadata={
                                        "headers": cleaned[0],
                                        "pdf_page_number": page_number,
                                    },
                                )
                            )
        except Exception as exc:
            # Text extraction is still useful when a table extractor cannot interpret the PDF.
            logger.warning(
                "pdf_table_extraction_failed",
                filename=path.name,
                error_type=type(exc).__name__,
            )

        if sparse_pages:
            try:
                import pytesseract
                from pdf2image import convert_from_path

                for page_number in sparse_pages:
                    images = convert_from_path(
                        path, first_page=page_number, last_page=page_number, dpi=250
                    )
                    if images:
                        text = pytesseract.image_to_string(images[0], lang=self.languages).strip()
                        if text:
                            page_label, text = extract_pdf_page_label(text, page_number)
                            elements.append(
                                ParsedElement(
                                    text=sanitize_document_text(text),
                                    content_type="image_ocr",
                                    page_number=page_label,
                                    section=f"Page {page_label}",
                                    image_id=f"pdf-page-{page_number}",
                                    metadata={"pdf_page_number": page_number},
                                )
                            )
            except Exception as exc:
                if not elements:
                    raise DocumentParseError(f"PDF requires OCR but OCR failed: {exc}") from exc

        metadata: dict[str, Any] = {
            "page_count": len(reader.pages),
            "author": str(reader.metadata.author)
            if reader.metadata and reader.metadata.author
            else None,
            "document_date": str(reader.metadata.creation_date)
            if reader.metadata and reader.metadata.creation_date
            else None,
        }
        return ParsedDocument(title, path.name, "pdf", mime_type, elements, metadata)


class ParserRegistry:
    def __init__(self, ocr_languages: str = "eng") -> None:
        parsers: list[DocumentParser] = [
            PDFParser(ocr_languages),
            DOCXParser(ocr_languages),
            CSVParser(),
            JSONParser(),
            HTMLParser(),
            ImageParser(ocr_languages),
            TextParser(),
        ]
        self._by_extension = {
            extension: parser for parser in parsers for extension in parser.extensions
        }

    def parse(self, path: Path, mime_type: str, source_url: str | None = None) -> ParsedDocument:
        parser = self._by_extension.get(path.suffix.lower())
        if not parser:
            raise UnsupportedDocumentError(f"No parser registered for {path.suffix}")
        parsed = parser.parse(path, mime_type, source_url)
        if not parsed.elements:
            raise DocumentParseError("Document contained no extractable text or tables")
        return parsed
