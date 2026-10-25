from __future__ import annotations

import json
from pathlib import Path

import pytest

from app.core.exceptions import DocumentParseError
from app.ingestion.parsers import ParserRegistry, table_to_markdown


def test_text_parser_preserves_headings(tmp_path: Path) -> None:
    path = tmp_path / "policy.txt"
    path.write_text("# Refund Policy\n\nRefunds are available for 30 days.\n", encoding="utf-8")
    parsed = ParserRegistry().parse(path, "text/plain")
    assert parsed.title == "policy"
    assert parsed.elements[0].content_type == "heading"
    assert parsed.elements[1].section == "Refund Policy"


def test_csv_parser_preserves_headers_and_rows(tmp_path: Path) -> None:
    path = tmp_path / "items.csv"
    path.write_text("sku,name,price\nA-1,Widget,10\n", encoding="utf-8")
    parsed = ParserRegistry().parse(path, "text/csv")
    assert parsed.elements[0].content_type == "table"
    assert "| sku | name | price |" in parsed.elements[0].text
    assert "| A-1 | Widget | 10 |" in parsed.elements[0].text


def test_json_parser_preserves_nested_paths_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "data.json"
    path.write_text(
        json.dumps({"policy": {"period": 30}, "items": [{"sku": "A"}, {"sku": "B"}]}),
        encoding="utf-8",
    )
    parsed = ParserRegistry().parse(path, "application/json")
    assert any(item.metadata.get("json_path") == "policy.period" for item in parsed.elements)
    assert any(item.content_type == "table" and item.section == "items" for item in parsed.elements)


def test_html_parser_keeps_sections_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "page.html"
    path.write_text(
        "<html><head><title>Guide</title></head><body><article><h2>Limits</h2>"
        "<p>The limit is ten.</p><table><tr><th>Name</th><th>Value</th></tr>"
        "<tr><td>Limit</td><td>10</td></tr></table></article></body></html>",
        encoding="utf-8",
    )
    parsed = ParserRegistry().parse(path, "text/html")
    assert parsed.title == "Guide"
    assert any(item.section == "Limits" for item in parsed.elements)
    assert any(item.content_type == "table" for item in parsed.elements)


def test_docx_parser_keeps_headings_and_table(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "guide.docx"
    document = Document()
    document.add_heading("Support Guide", 0)
    document.add_heading("Contact", 1)
    document.add_paragraph("Contact support through the portal.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Channel"
    table.cell(0, 1).text = "Hours"
    table.cell(1, 0).text = "Portal"
    table.cell(1, 1).text = "24/7"
    document.save(path)
    parsed = ParserRegistry().parse(
        path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert parsed.title == "Support Guide"
    assert any(item.section == "Contact" for item in parsed.elements)
    assert any(item.content_type == "table" for item in parsed.elements)


def test_corrupted_json_fails_safely(tmp_path: Path) -> None:
    path = tmp_path / "bad.json"
    path.write_text("{not-json", encoding="utf-8")
    with pytest.raises(DocumentParseError):
        ParserRegistry().parse(path, "application/json")


def test_table_markdown_escapes_cell_pipes() -> None:
    assert "a\\|b" in table_to_markdown(["Name"], [["a|b"]])
